"""Version Excel de la table de correspondance, pour que la relectrice puisse y répondre case par case.

Produit un classeur à trois feuilles : « Lisez-moi » (mode d'emploi et couleurs), « Questions » (les
cinq points à trancher, avec une colonne de réponse) et « Table » (une ligne par instrument, avec
deux colonnes à remplir : « Votre avis », en liste déroulante, et « Votre correction »).

Les clés bibliographiques deviennent des intitulés de texte, les chapitres des liens vers le site.
Les couleurs reprennent celles de la page : orange (proposition, à confirmer), jaune (piste à
vérifier), bleu (case vide).

    uv run --with openpyxl python scripts/correspondance_xlsx.py sortie.xlsx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import correspondance as c

SITE = "https://benjello.github.io/precis-socio-fiscal-tunisie/fr/"

COLONNES = [
    ("Famille de droit", 22),
    ("Instrument", 30),
    ("Catégorie", 16),
    ("Institution gestionnaire", 28),
    ("Texte de référence", 42),
    ("Principaux modificatifs", 42),
    ("Texte d'application", 42),
    ("Statut du fondateur", 30),
    ("Chapitre du précis", 30),
    ("Votre avis", 16),
    ("Votre correction ou le texte à ajouter", 46),
]

QUESTIONS = [
    "Programme national d'aide aux familles nécessiteuses : les circulaires de 1986, 1993 et 2011 sont-elles "
    "bien rangées (1986 : texte de référence ; 1993 et 2011 : modificatifs) ? Quels sont le numéro et le jour "
    "de la circulaire de mai 2011 ?",
    "Fonctions administratives et voies de recours : que recouvre pour vous la distinction des fonctions "
    "administratives ? Faut-il une colonne « Voies de recours », ou la juridiction compétente suffit-elle ?",
    "Institutions proposées (cases orange, colonne « Institution gestionnaire ») : impôt sur les sociétés, "
    "droits de consommation, code du travail, statut général. Faut-il confirmer, corriger ou écarter ?",
    "Pistes de textes (cases jaunes) : décret de 2009 sur l'admission temporaire et recueil du code des "
    "douanes ; décret de 2016, décret-loi de 2024 et loi de 1997 pour les salariés agricoles. Sont-ils les "
    "bons, et quel est le texte d'application principal ?",
    "Textes d'application manquants (cases bleues, colonne « Texte d'application ») : connaissez-vous le "
    "texte principal pour l'impôt sur le revenu, l'impôt sur les sociétés, l'impôt sur la fortune, les "
    "pensions des non-salariés, le code du travail, l'enseignement supérieur privé, les cantines scolaires "
    "et le prêt universitaire des caisses ?",
]


def intitule(refs: dict, cle: str, locator: str | None) -> str:
    x = refs.get(cle, {})
    titre = x.get("title", cle)
    return f"{titre}, {locator}" if locator else titre


def texte(refs: dict, valeur: str) -> str:
    """Une ligne par texte ; le balisage italique du CSV est retiré."""
    lignes = []
    for cle, locator, est_texte in c.elements(valeur):
        lignes.append(cle.replace("*", "") if est_texte else intitule(refs, cle, locator))
    return "\n".join(lignes)


def statut(refs: dict, valeur: str) -> str:
    """Le statut est du Markdown : les citations [@clé, locator] deviennent des intitulés."""
    if not valeur.strip():
        return "—"
    def remplace(m):
        cle, _, locator = m.group(1).partition(",")
        return "(" + intitule(refs, cle.strip(), locator.strip() or None) + ")"
    return re.sub(r"\[@([^\]]+)\]", remplace, valeur).replace("*", "")


def couleur(brut: str) -> str | None:
    if not brut.strip():
        return "CFE3FF"
    if "*Proposition" in brut or "*À confirmer" in brut:
        return "FFD9A8"
    if "*Piste" in brut:
        return "FFF3A8"
    return None


def classeur(chemin: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    refs = c.references()
    wb = Workbook()

    aide = wb.active
    aide.title = "Lisez-moi"
    lignes = [
        ("Table de correspondance : instruments, institutions et textes", True),
        ("Version de travail du 30 octobre 2026, non publiée. Merci de votre relecture.", False),
        ("", False),
        ("Comment répondre", True),
        ("1. Feuille « Questions » : cinq points sur lesquels votre avis m'est le plus utile. Écrivez dans la colonne C.", False),
        ("2. Feuille « Table » : une ligne par instrument. Colonne « Votre avis » : choisissez dans la liste. "
         "Colonne « Votre correction » : écrivez le texte juste ou ce qu'il faut ajouter. Quelques mots suffisent.", False),
        ("3. Vous pouvez écrire en français ou en arabe. Renvoyez le fichier par mail, ou répondez dans le corps du mail.", False),
        ("", False),
        ("Couleurs des cases", True),
        ("Orange : proposition de ma part, ou point à confirmer.", False),
        ("Jaune : piste d'une recherche en ligne, textes non consultés au Journal officiel, à vérifier.", False),
        ("Bleu : case vide, texte à établir.", False),
        ("Sans couleur : relevé dans les chapitres du précis, avec sa source.", False),
    ]
    for i, (t, gras) in enumerate(lignes, 1):
        cell = aide.cell(row=i, column=1, value=t)
        cell.font = Font(bold=gras, size=13 if i == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for ligne, hexa in ((10, "FFD9A8"), (11, "FFF3A8"), (12, "CFE3FF")):
        aide.cell(row=ligne, column=1).fill = PatternFill("solid", fgColor=hexa)
    aide.column_dimensions["A"].width = 110

    q = wb.create_sheet("Questions")
    for j, (titre, largeur) in enumerate((("N°", 5), ("Question", 100), ("Votre réponse", 60)), 1):
        cell = q.cell(row=1, column=j, value=titre)
        cell.font = Font(bold=True)
        q.column_dimensions[chr(64 + j)].width = largeur
    for i, texte_q in enumerate(QUESTIONS, 2):
        q.cell(row=i, column=1, value=i - 1)
        q.cell(row=i, column=2, value=texte_q).alignment = Alignment(wrap_text=True, vertical="top")
        q.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        q.cell(row=i, column=3).fill = PatternFill("solid", fgColor="FFF7EC")

    t = wb.create_sheet("Table")
    for j, (titre, largeur) in enumerate(COLONNES, 1):
        cell = t.cell(row=1, column=j, value=titre)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        t.column_dimensions[chr(64 + j)].width = largeur
    t.freeze_panes = "C2"
    validation = DataValidation(type="list", formula1='"Correct,À corriger,À compléter,Je ne sais pas"', allow_blank=True)
    t.add_data_validation(validation)

    for i, ligne in enumerate(c.lire(), 2):
        valeurs = [
            c.FAMILLES[ligne["famille"]],
            ligne["instrument"],
            ligne["categorie"],
            ligne["institution"].replace("*", "") or "à établir",
            texte(refs, ligne["fondateur"]),
            texte(refs, ligne["modificatifs"]) or "—",
            texte(refs, ligne["application"]) or "—",
            statut(refs, ligne["statut"]),
        ]
        for j, v in enumerate(valeurs, 1):
            t.cell(row=i, column=j, value=v)
        premiers = c.chapitres(ligne["chapitres"])
        if premiers:
            libelle, cible = premiers[0]
            cell = t.cell(row=i, column=9, value=libelle)
            cell.hyperlink = SITE + cible
            cell.font = Font(color="1A3D7A", underline="single")
        for col, brut in ((4, ligne["institution"]), (5, ligne["fondateur"]), (6, ligne["modificatifs"]), (7, ligne["application"])):
            hexa = couleur(brut)
            if hexa:
                t.cell(row=i, column=col).fill = PatternFill("solid", fgColor=hexa)
        validation.add(t.cell(row=i, column=10))
        t.cell(row=i, column=10).fill = PatternFill("solid", fgColor="FFF7EC")
        t.cell(row=i, column=11).fill = PatternFill("solid", fgColor="FFF7EC")
        for j in range(1, len(COLONNES) + 1):
            t.cell(row=i, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    t.auto_filter.ref = f"A1:{chr(64 + len(COLONNES))}{t.max_row}"
    wb.save(chemin)


if __name__ == "__main__":
    classeur(Path(sys.argv[1]))
